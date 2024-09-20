from openpyxl import load_workbook
from openpyxl.drawing.image import Image
import xlsxwriter
import os


def insert():
    base_dir = "/Users/hiephuynh/Downloads"
    excel_path = "image.xlsx"
    # image_path = os.path.join(base_dir, "images/image.png")
    # Load the workbook and select the worksheet
    workbook = load_workbook("venv/image.xlsx")
    # workbook = xlsxwriter.Workbook("venv/image.xlsx")
    # sheet = workbook.get_worksheet_by_name("image")
    sheet = workbook["image"]

    path_list={}
    count =1
    # for file in base_dir.iterdir():
    #     print(file)
    #     count+=1
    #     if count==10:
    #         break
    for date,symbol in sheet.iter_rows(values_only=True):
        if date:
            date_str=date.strftime('%-m:%-d:%-y')
            file_name =date_str+" "+symbol+".PNG"
            path=os.path.join(base_dir,file_name)
            image= Image(path)
            cell="C"+str(count)
            sheet.add_image(image,cell)
            count+=1
            if image:
                print(image)
        else:
            print(date)
            break

    
    workbook.save("venv/image.xlsx")
    print("saved")



def main():
    insert()


if __name__=="__main__":
    main()

